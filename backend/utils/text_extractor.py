"""
Text extraction utilities for various document formats.
"""

import os
import logging
from typing import Optional
from docx import Document
import mistune
import html

logger = logging.getLogger(__name__)


class TextExtractor:
    """Extract text content from various document formats."""
    
    @staticmethod
    def extract_text(file_path: str) -> str:
        """
        Extract text from a file based on its extension.
        
        Args:
            file_path: Path to the file
            
        Returns:
            Extracted text content
            
        Raises:
            ValueError: If file format is not supported
            Exception: If extraction fails
        """
        file_ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if file_ext == '.txt':
                return TextExtractor._extract_from_txt(file_path)
            elif file_ext == '.docx':
                return TextExtractor._extract_from_docx(file_path)
            elif file_ext == '.html':
                return TextExtractor._extract_from_html(file_path)
            elif file_ext == '.md':
                return TextExtractor._extract_from_markdown(file_path)
            elif file_ext == '.pdf':
                # PDF extraction would require PyMuPDF or similar
                # For now, we'll raise an error for PDFs
                raise ValueError("PDF extraction not yet implemented. Please use .txt, .docx, .html, or .md files.")
            else:
                raise ValueError(f"Unsupported file format: {file_ext}")
        except Exception as e:
            logger.error(f"Error extracting text from {file_path}: {str(e)}")
            raise
    
    @staticmethod
    def _extract_from_txt(file_path: str) -> str:
        """Extract text from a plain text file."""
        encodings = ['utf-8', 'utf-16', 'latin-1', 'cp1252']
        
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        
        # If all encodings fail, try with errors='ignore'
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    
    @staticmethod
    def _extract_from_docx(file_path: str) -> str:
        """Extract text from a DOCX file."""
        doc = Document(file_path)
        
        # Extract text from all paragraphs
        paragraphs = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                paragraphs.append(text)
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text and text not in paragraphs:
                        paragraphs.append(text)
        
        return '\n\n'.join(paragraphs)
    
    @staticmethod
    def _extract_from_html(file_path: str) -> str:
        """Extract text from an HTML file."""
        with open(file_path, 'r', encoding='utf-8') as f:
            html_content = f.read()
        
        # Simple HTML tag removal (for basic HTML)
        # For complex HTML, consider using BeautifulSoup
        import re
        
        # Remove script and style elements
        html_content = re.sub(r'<script.*?</script>', '', html_content, flags=re.DOTALL)
        html_content = re.sub(r'<style.*?</style>', '', html_content, flags=re.DOTALL)
        
        # Remove HTML tags
        text = re.sub(r'<[^>]+>', ' ', html_content)
        
        # Unescape HTML entities
        text = html.unescape(text)
        
        # Clean up whitespace
        text = ' '.join(text.split())
        
        return text
    
    @staticmethod
    def _extract_from_markdown(file_path: str) -> str:
        """Extract text from a Markdown file."""
        with open(file_path, 'r', encoding='utf-8') as f:
            md_content = f.read()
        
        # Convert markdown to HTML, then extract text
        markdown = mistune.create_markdown()
        html_content = markdown(md_content)
        
        # Remove HTML tags
        import re
        text = re.sub(r'<[^>]+>', ' ', html_content)
        text = html.unescape(text)
        text = ' '.join(text.split())
        
        return text
    
    @staticmethod
    def save_text_to_file(text: str, output_path: str, format: str = 'txt'):
        """
        Save text content to a file in the specified format.
        
        Args:
            text: Text content to save
            output_path: Output file path (without extension)
            format: Output format ('txt', 'docx')
        """
        if format == 'txt':
            with open(f"{output_path}.txt", 'w', encoding='utf-8') as f:
                f.write(text)
        elif format == 'docx':
            doc = Document()
            # Split text into paragraphs
            paragraphs = text.split('\n\n')
            for para in paragraphs:
                if para.strip():
                    doc.add_paragraph(para.strip())
            doc.save(f"{output_path}.docx")
        else:
            raise ValueError(f"Unsupported output format: {format}")